# Copyright (c) Huawei Technologies Co., Ltd. 2026. All rights reserved.

"""Tests for the doc_converter parser registry and the pdf/docx/pptx/xlsx/xls/_table parsers."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest

from chrys.service.tools.builtins.doc_converter.artifacts import (
    MAX_IMAGE_OCCURRENCES,
    DocumentImageSink,
)
from tests.service.tools._doc_converter_fakes import (
    _png_bytes,
    _session_artifact_path,
)


class _FakePdfImages:
    """Stand-in for pypdf's ``page.images``: ``keys()`` lists every candidate, ``[]`` builds one image object per lookup."""

    def __init__(self, entries: dict[str, dict[str, object]]) -> None:
        self._entries = entries
        self.keys_materialized = 0
        self.retrieval_count = 0

    def keys(self) -> list[str]:
        self.keys_materialized += len(self._entries)
        return list(self._entries)

    def __getitem__(self, key: str) -> MagicMock:
        self.retrieval_count += 1
        return MagicMock(**self._entries[key])


def _pdf_image(
    data: bytes, name: str, *, is_displayed: bool = True, indirect_reference: object = None
) -> dict[str, object]:
    return {"data": data, "name": name, "is_displayed": is_displayed, "indirect_reference": indirect_reference}


# ---------------------------------------------------------------------------
# Parser registry
# ---------------------------------------------------------------------------


def test_parser_registry_all_extensions() -> None:
    """All expected extensions are registered."""
    from chrys.service.tools.builtins.doc_converter.registry import supported_extensions

    exts = supported_extensions()
    assert ".pdf" in exts
    assert ".docx" in exts
    assert ".pptx" in exts
    assert ".xlsx" in exts
    assert ".xls" in exts
    assert ".epub" not in exts


def test_parser_registry_get_parser() -> None:
    """get_parser returns correct parser type for each extension."""
    from chrys.service.tools.builtins.doc_converter.parsers.docx import DocxParser
    from chrys.service.tools.builtins.doc_converter.parsers.pdf import PdfParser
    from chrys.service.tools.builtins.doc_converter.parsers.pptx import PptxParser
    from chrys.service.tools.builtins.doc_converter.parsers.xls import XlsParser
    from chrys.service.tools.builtins.doc_converter.parsers.xlsx import XlsxParser
    from chrys.service.tools.builtins.doc_converter.registry import get_parser

    assert isinstance(get_parser(".pdf"), PdfParser)
    assert isinstance(get_parser(".docx"), DocxParser)
    assert isinstance(get_parser(".pptx"), PptxParser)
    assert isinstance(get_parser(".xlsx"), XlsxParser)
    assert isinstance(get_parser(".xls"), XlsParser)
    assert get_parser(".epub") is None
    assert get_parser(".txt") is None


def test_parser_registry_case_insensitive() -> None:
    from chrys.service.tools.builtins.doc_converter.registry import get_parser

    assert get_parser(".PDF") is not None
    assert get_parser(".Docx") is not None


# ---------------------------------------------------------------------------
# Individual parser unit tests (mock the library imports)
# ---------------------------------------------------------------------------


def test_pdf_parser_output() -> None:
    """PdfParser produces page-based Markdown headings."""
    from chrys.service.tools.builtins.doc_converter.parsers.pdf import PdfParser

    mock_page1 = MagicMock()
    mock_page1.extract_text.return_value = "Hello world"
    mock_page2 = MagicMock()
    mock_page2.extract_text.return_value = "Second page content"

    mock_reader = MagicMock()
    mock_reader.pages = [mock_page1, mock_page2]

    mock_pypdf = MagicMock()
    mock_pypdf.PdfReader.return_value = mock_reader

    with patch.dict("sys.modules", {"pypdf": mock_pypdf}):
        parser = PdfParser()
        result = parser.parse("/fake.pdf")

    assert "# Page 1" in result.markdown
    assert "Hello world" in result.markdown
    assert "# Page 2" in result.markdown
    assert "Second page content" in result.markdown


def test_pdf_parser_extracts_image_only_page_and_keeps_no_text_placeholder(tmp_path: Path) -> None:
    from PIL import Image

    from chrys.service.tools.builtins.doc_converter.parsers.pdf import PdfParser

    pdf = tmp_path / "image.pdf"
    Image.new("RGB", (8, 8), (255, 0, 0)).save(pdf, format="PDF")
    sink = DocumentImageSink(tmp_path / "session" / "doc_converter", source_stem="image")

    result = PdfParser().parse(str(pdf), image_sink=sink)

    assert "# Page 1\n\n(no text content)" in result.markdown
    assert len(result.visuals) == 1
    assert _session_artifact_path(tmp_path / "session", result.visuals[0].reference).exists()


def test_pdf_parser_deduplicates_repeated_indirect_reference(tmp_path: Path) -> None:
    from chrys.service.tools.builtins.doc_converter.parsers.pdf import PdfParser

    image_data = _png_bytes()
    indirect_reference = object()

    pages = []
    for _ in range(2):
        page = MagicMock()
        page.extract_text.return_value = "text"
        page.images = _FakePdfImages(
            {"/Im0": _pdf_image(image_data, "logo.png", indirect_reference=indirect_reference)}
        )
        pages.append(page)
    mock_pypdf = MagicMock()
    mock_pypdf.PdfReader.return_value.pages = pages
    sink = DocumentImageSink(tmp_path / "session" / "doc_converter", source_stem="report")

    with (
        patch.dict("sys.modules", {"pypdf": mock_pypdf}),
        patch.object(sink, "save_image", wraps=sink.save_image) as save_image,
    ):
        result = PdfParser().parse("/fake.pdf", image_sink=sink)

    assert len(result.visuals) == 2
    assert result.visuals[0].reference == result.visuals[1].reference
    assert save_image.call_count == 1
    assert len(list((tmp_path / "session" / "doc_converter").iterdir())) == 1


def test_pdf_parser_skips_undisplayed_resource_without_writing(tmp_path: Path) -> None:
    from chrys.service.tools.builtins.doc_converter.parsers.pdf import PdfParser

    page = MagicMock()
    page.extract_text.return_value = "text"
    page.images = _FakePdfImages(
        {"/Unused": _pdf_image(_png_bytes(), "unused.png", is_displayed=False, indirect_reference=object())}
    )
    mock_pypdf = MagicMock()
    mock_pypdf.PdfReader.return_value.pages = [page]
    sink = DocumentImageSink(tmp_path / "session" / "doc_converter", source_stem="report")

    with patch.dict("sys.modules", {"pypdf": mock_pypdf}):
        result = PdfParser().parse("/fake.pdf", image_sink=sink)

    assert result.visuals == ()
    assert not (tmp_path / "session" / "doc_converter").exists()


def test_pdf_parser_corrupt_image_preserves_text_and_other_images(tmp_path: Path) -> None:
    from chrys.service.tools.builtins.doc_converter.parsers.pdf import PdfParser

    page = MagicMock()
    page.extract_text.return_value = "Preserved page text"
    page.images = _FakePdfImages(
        {
            "/Corrupt": _pdf_image(b"not-an-image", "/Corrupt.png"),
            "/Valid": _pdf_image(_png_bytes(), "/Valid.png"),
        }
    )
    mock_pypdf = MagicMock()
    mock_pypdf.PdfReader.return_value.pages = [page]
    sink = DocumentImageSink(tmp_path / "session" / "doc_converter", source_stem="report")

    with patch.dict("sys.modules", {"pypdf": mock_pypdf}):
        result = PdfParser().parse("/fake.pdf", image_sink=sink)

    assert "Preserved page text" in result.markdown
    assert len(result.visuals) == 1
    assert len(result.warnings) == 1
    assert "could not be decoded or normalized" in result.warnings[0]


@pytest.mark.parametrize(
    ("key_template", "image_name"),
    [("/Im{index}", "logo.png"), ("~{index}~", "inline.png")],
    ids=["xobject", "inline"],
)
def test_pdf_parser_caps_image_retrievals_after_materializing_keys(
    tmp_path: Path, key_template: str, image_name: str
) -> None:
    """``keys()`` materializes every candidate (pypdf scans inline images eagerly); retrievals stop at the cap."""
    from chrys.service.tools.builtins.doc_converter.parsers.pdf import PdfParser

    candidate_count = MAX_IMAGE_OCCURRENCES + 2
    images = _FakePdfImages(
        {key_template.format(index=index): _pdf_image(_png_bytes(), image_name) for index in range(candidate_count)}
    )
    page = MagicMock()
    page.extract_text.return_value = "text"
    page.images = images
    mock_pypdf = MagicMock()
    mock_pypdf.PdfReader.return_value.pages = [page]
    sink = DocumentImageSink(tmp_path / "session" / "doc_converter", source_stem="report")

    with patch.dict("sys.modules", {"pypdf": mock_pypdf}):
        result = PdfParser().parse("/fake.pdf", image_sink=sink)

    assert images.keys_materialized == candidate_count
    assert images.retrieval_count == MAX_IMAGE_OCCURRENCES
    assert len(result.visuals) == MAX_IMAGE_OCCURRENCES
    assert len(result.warnings) == 1
    assert "2 image candidate(s)" in result.warnings[0]
    assert sink.warnings == result.warnings


def test_docx_parser_headings_and_tables() -> None:
    """DocxParser maps heading styles to Markdown headings and extracts tables in order."""
    from chrys.service.tools.builtins.doc_converter.parsers.docx import DocxParser

    # Real classes so isinstance() checks work with the mocked docx.table.Table
    class _FakeDocxTable:
        def __init__(self, row_data: list[list[str]]):
            self.rows = [MagicMock(cells=[MagicMock(text=c) for c in r]) for r in row_data]

    mock_para1 = MagicMock()
    mock_para1.text = "My Title"
    mock_para1.style.name = "Heading 1"

    mock_para2 = MagicMock()
    mock_para2.text = "Some body text"
    mock_para2.style.name = "Normal"

    mock_table = _FakeDocxTable([["Name", "Age"], ["Alice", "30"]])

    mock_para3 = MagicMock()
    mock_para3.text = "Subsection"
    mock_para3.style.name = "Heading 2"

    mock_doc = MagicMock()
    mock_doc.iter_inner_content.return_value = iter([mock_para1, mock_para2, mock_table, mock_para3])

    mock_docx = MagicMock()
    mock_docx.Document.return_value = mock_doc
    mock_table_mod = MagicMock()
    mock_table_mod.Table = _FakeDocxTable

    with patch.dict("sys.modules", {"docx": mock_docx, "docx.table": mock_table_mod}):
        parser = DocxParser()
        result = parser.parse("/fake.docx")

    assert "# My Title" in result.markdown
    assert "Some body text" in result.markdown
    assert "| Name | Age |" in result.markdown
    assert "| Alice | 30 |" in result.markdown
    assert "## Subsection" in result.markdown
    # Verify ordering: title before table before subsection
    assert result.markdown.index("My Title") < result.markdown.index("Name | Age") < result.markdown.index("Subsection")


def test_docx_parser_extracts_package_wide_header_image(tmp_path: Path) -> None:
    from docx import Document

    from chrys.service.tools.builtins.doc_converter.parsers.docx import DocxParser

    document = Document()
    document.add_paragraph("Body text")
    document.sections[0].header.paragraphs[0].add_run().add_picture(BytesIO(_png_bytes()))
    path = tmp_path / "header.docx"
    document.save(path)
    sink = DocumentImageSink(tmp_path / "session" / "doc_converter", source_stem="header")

    result = DocxParser().parse(str(path), image_sink=sink)

    assert result.markdown == "Body text"
    assert len(result.visuals) == 1
    assert result.visuals[0].location == "Document"
    assert _session_artifact_path(tmp_path / "session", result.visuals[0].reference).exists()


def test_docx_parser_package_image_parts_deduplicate_reused_blob(tmp_path: Path) -> None:
    from docx import Document

    from chrys.service.tools.builtins.doc_converter.parsers.docx import DocxParser

    image = BytesIO(_png_bytes())
    document = Document()
    document.add_picture(image)
    image.seek(0)
    document.add_picture(image)
    path = tmp_path / "reused.docx"
    document.save(path)
    sink = DocumentImageSink(tmp_path / "session" / "doc_converter", source_stem="reused")

    result = DocxParser().parse(str(path), image_sink=sink)

    assert result.markdown == ""
    assert len(result.visuals) == 1
    assert len(list((tmp_path / "session" / "doc_converter").iterdir())) == 1


def test_pptx_parser_slides() -> None:
    """PptxParser produces slide-based Markdown headings."""
    from chrys.service.tools.builtins.doc_converter.parsers.pptx import PptxParser

    mock_title = MagicMock()
    mock_title.text = "Intro Slide"

    mock_shape = MagicMock()
    mock_shape.has_text_frame = True
    mock_shape.has_table = False
    mock_para = MagicMock()
    mock_para.text = "Bullet point"
    mock_shape.text_frame.paragraphs = [mock_para]

    mock_slide = MagicMock()
    mock_slide.shapes.title = mock_title
    mock_slide.shapes.__iter__ = MagicMock(return_value=iter([mock_shape]))

    mock_prs = MagicMock()
    mock_prs.slides = [mock_slide]

    mock_pptx = MagicMock()
    mock_pptx.Presentation.return_value = mock_prs

    with patch.dict("sys.modules", {"pptx": mock_pptx}):
        parser = PptxParser()
        result = parser.parse("/fake.pptx")

    assert "# Slide 1: Intro Slide" in result.markdown
    assert "Bullet point" in result.markdown


def test_pptx_parser_table_shape() -> None:
    """PptxParser extracts tables from table shapes."""
    from chrys.service.tools.builtins.doc_converter.parsers.pptx import PptxParser

    # Text shape
    text_shape = MagicMock()
    text_shape.has_table = False
    text_shape.has_text_frame = True
    text_shape.text_frame.paragraphs = [MagicMock(text="Intro text")]

    # Table shape
    table_shape = MagicMock()
    table_shape.has_table = True
    table_shape.has_text_frame = False
    table_shape.table.rows = [
        MagicMock(cells=[MagicMock(text="Col A"), MagicMock(text="Col B")]),
        MagicMock(cells=[MagicMock(text="val1"), MagicMock(text="val2")]),
    ]

    mock_slide = MagicMock()
    mock_slide.shapes.title = None
    mock_slide.shapes.__iter__ = MagicMock(return_value=iter([text_shape, table_shape]))

    mock_prs = MagicMock()
    mock_prs.slides = [mock_slide]

    mock_pptx = MagicMock()
    mock_pptx.Presentation.return_value = mock_prs

    with patch.dict("sys.modules", {"pptx": mock_pptx}):
        parser = PptxParser()
        result = parser.parse("/fake.pptx")

    assert "# Slide 1" in result.markdown
    assert "Intro text" in result.markdown
    assert "| Col A | Col B |" in result.markdown
    assert "| val1 | val2 |" in result.markdown


def test_pptx_parser_extracts_top_level_and_grouped_pictures(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    from chrys.service.tools.builtins.doc_converter.parsers.pptx import PptxParser

    image = tmp_path / "pixel.png"
    image.write_bytes(_png_bytes())
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    top_level = slide.shapes.add_picture(str(image), Inches(1), Inches(1))
    slide.shapes.add_group_shape([top_level])
    deck = tmp_path / "deck.pptx"
    presentation.save(deck)
    sink = DocumentImageSink(tmp_path / "session" / "doc_converter", source_stem="deck")

    result = PptxParser().parse(str(deck), image_sink=sink)

    assert len(result.visuals) == 1
    assert result.visuals[0].location == "Slide 1"
    assert _session_artifact_path(tmp_path / "session", result.visuals[0].reference).exists()


def test_pptx_parser_extracts_populated_picture_placeholder(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.enum.shapes import PP_PLACEHOLDER
    from pptx.shapes.picture import Picture

    from chrys.service.tools.builtins.doc_converter.parsers.pptx import PptxParser

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[8])
    placeholder = next(
        candidate for candidate in slide.placeholders if candidate.placeholder_format.type == PP_PLACEHOLDER.PICTURE
    )
    populated = placeholder.insert_picture(BytesIO(_png_bytes()))
    assert isinstance(populated, Picture)
    assert populated.shape_type.name == "PLACEHOLDER"
    deck = tmp_path / "placeholder.pptx"
    presentation.save(deck)
    sink = DocumentImageSink(tmp_path / "session" / "doc_converter", source_stem="placeholder")

    result = PptxParser().parse(str(deck), image_sink=sink)

    assert len(result.visuals) == 1
    assert _session_artifact_path(tmp_path / "session", result.visuals[0].reference).exists()


def test_pptx_parser_repeated_logo_across_slides_stores_once(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    from chrys.service.tools.builtins.doc_converter.parsers.pptx import PptxParser

    image = tmp_path / "logo.png"
    image.write_bytes(_png_bytes())
    presentation = Presentation()
    for _ in range(2):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_picture(str(image), Inches(1), Inches(1))
    deck = tmp_path / "repeated-logo.pptx"
    presentation.save(deck)
    sink = DocumentImageSink(tmp_path / "session" / "doc_converter", source_stem="repeated-logo")

    result = PptxParser().parse(str(deck), image_sink=sink)

    assert len(result.visuals) == 2
    assert result.visuals[0].reference == result.visuals[1].reference
    assert len(list((tmp_path / "session" / "doc_converter").iterdir())) == 1


def test_pptx_parser_skips_linked_only_picture_with_bounded_warning(tmp_path: Path) -> None:
    from pptx.shapes.picture import Picture

    from chrys.service.tools.builtins.doc_converter.parsers.pptx import PptxParser

    class _LinkedPicture(Picture):
        @property
        def has_table(self) -> bool:
            return False

        @property
        def has_text_frame(self) -> bool:
            return False

        @property
        def image(self):
            raise ValueError("no embedded image")

    linked = object.__new__(_LinkedPicture)
    slide = MagicMock()
    slide.shapes.title = None
    slide.shapes.__iter__ = MagicMock(side_effect=lambda: iter([linked]))
    presentation = MagicMock()
    presentation.slides = [slide]
    sink = DocumentImageSink(tmp_path / "session" / "doc_converter", source_stem="linked")

    with patch("pptx.Presentation", return_value=presentation):
        result = PptxParser().parse("/fake.pptx", image_sink=sink)

    assert result.visuals == ()
    assert len(result.warnings) == 1
    assert "linked image(s)" in result.warnings[0]


def test_xlsx_parser_tables() -> None:
    """XlsxParser produces sheet-based Markdown tables."""
    from chrys.service.tools.builtins.doc_converter.parsers.xlsx import XlsxParser

    mock_ws = MagicMock()
    mock_ws.iter_rows.return_value = [("Name", "Age"), ("Alice", 30), ("Bob", 25)]

    mock_wb = MagicMock()
    mock_wb.sheetnames = ["Sheet1"]
    mock_wb.__getitem__ = lambda self, key: mock_ws

    mock_openpyxl = MagicMock()
    mock_openpyxl.load_workbook.return_value = mock_wb

    with patch.dict("sys.modules", {"openpyxl": mock_openpyxl}):
        parser = XlsxParser()
        result = parser.parse("/fake.xlsx", image_sink=MagicMock(spec=DocumentImageSink))

    assert "# Sheet: Sheet1" in result.markdown
    assert "| Name | Age |" in result.markdown
    assert "| Alice | 30 |" in result.markdown
    assert result.warnings == ()


def test_xlsx_parser_with_embedded_image_stays_text_only_without_result_warning(tmp_path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as SpreadsheetImage

    from chrys.service.tools.builtins.doc_converter.parsers.xlsx import XlsxParser

    image_path = tmp_path / "chart.png"
    image_path.write_bytes(_png_bytes())
    workbook = Workbook()
    worksheet = workbook.active
    worksheet["A1"] = "Report"
    worksheet.add_image(SpreadsheetImage(image_path), "B2")
    spreadsheet = tmp_path / "report.xlsx"
    workbook.save(spreadsheet)
    workbook.close()
    sink = DocumentImageSink(tmp_path / "session" / "doc_converter", source_stem="report")

    result = XlsxParser().parse(str(spreadsheet), image_sink=sink)

    assert result.visuals == ()
    assert result.warnings == ()
    assert not (tmp_path / "session" / "doc_converter").exists()


def test_xls_parser_tables() -> None:
    """XlsParser produces sheet-based Markdown tables."""
    from chrys.service.tools.builtins.doc_converter.parsers.xls import XlsParser

    mock_sheet = MagicMock()
    mock_sheet.name = "Data"
    mock_sheet.nrows = 2
    mock_sheet.ncols = 2
    mock_sheet.cell_value = lambda r, c: [["Header1", "Header2"], ["val1", "val2"]][r][c]

    mock_wb = MagicMock()
    mock_wb.sheets.return_value = [mock_sheet]

    mock_xlrd = MagicMock()
    mock_xlrd.open_workbook.return_value = mock_wb

    with patch.dict("sys.modules", {"xlrd": mock_xlrd}):
        parser = XlsParser()
        result = parser.parse("/fake.xls", image_sink=MagicMock(spec=DocumentImageSink))

    assert "# Sheet: Data" in result.markdown
    assert "| Header1 | Header2 |" in result.markdown
    assert "| val1 | val2 |" in result.markdown
    assert result.warnings == ()


def test_protocol_compliance() -> None:
    """All parsers satisfy the DocParser protocol."""
    from chrys.service.tools.builtins.doc_converter.parsers.base import DocParser
    from chrys.service.tools.builtins.doc_converter.parsers.docx import DocxParser
    from chrys.service.tools.builtins.doc_converter.parsers.pdf import PdfParser
    from chrys.service.tools.builtins.doc_converter.parsers.pptx import PptxParser
    from chrys.service.tools.builtins.doc_converter.parsers.xls import XlsParser
    from chrys.service.tools.builtins.doc_converter.parsers.xlsx import XlsxParser

    for cls in [PdfParser, DocxParser, PptxParser, XlsxParser, XlsParser]:
        assert isinstance(cls(), DocParser), f"{cls.__name__} does not satisfy DocParser protocol"


# ---------------------------------------------------------------------------
# Markdown table pipe escaping
# ---------------------------------------------------------------------------


def test_table_pipe_escaping() -> None:
    """Pipe characters in cell values are escaped so they don't break the table."""
    from chrys.service.tools.builtins.doc_converter.parsers._table import rows_to_markdown_table

    rows = [("Header", "Formula"), ("A|B", "x | y")]
    result = rows_to_markdown_table(rows)
    assert "| Header | Formula |" in result
    assert r"| A\|B | x \| y |" in result
